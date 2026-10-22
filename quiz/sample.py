"""File Word mẫu cho giáo viên tải về, sửa nội dung rồi tải lên."""
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

GUIDE = [
    'Quy tắc soạn đề (đọc xong có thể xóa phần hướng dẫn này):',
    '- Mỗi câu bắt đầu bằng "Câu 1:", "Câu 2:", ... Nội dung câu hỏi được viết nhiều dòng.',
    '- Mỗi phương án bắt đầu bằng "A.", "B.", "C.", "D." (có thể có thêm E., F.). Nhiều phương án trên một dòng thì cách nhau bằng phím Tab.',
    '- Sau các phương án là dòng "Đáp án: B". Câu có nhiều đáp án đúng: "Đáp án: A, C".',
    '- Lời giải hoặc ghi chú viết sau dòng "Đáp án" sẽ không hiện cho sinh viên.',
    '- Không cần đánh số liên tục: hệ thống rút ngẫu nhiên câu từ ngân hàng cho từng sinh viên.',
]
QUESTIONS = [
    ('Thủ đô của Việt Nam là thành phố nào?', ['Hà Nội', 'Huế', 'Đà Nẵng', 'TP. Hồ Chí Minh'], 'A', None),
    ('Kết quả của phép tính 2 + 3 × 4 là bao nhiêu?', ['20', '14', '24', '9'], 'B', 'Lời giải: nhân trước, cộng sau: 2 + 12 = 14.'),
    ('Những số nào sau đây là số nguyên tố?', ['2', '9', '7', '15'], 'A, C', None),
    ('Trong Python, kiểu dữ liệu nào là bất biến (immutable)?', ['list', 'tuple', 'dict', 'str'], 'B, D', None),
    ('Cho đoạn mã:\nx = [1, 2, 3]\nprint(len(x))\nKết quả in ra là gì?', ['1', '2', '3', 'Báo lỗi'], 'C', None),
]


def sample_docx():
    doc = Document()
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(13)
    doc.add_heading('NGÂN HÀNG CÂU HỎI MẪU', level=1)
    for line in GUIDE:
        run = doc.add_paragraph().add_run(line)
        run.italic, run.font.color.rgb = True, RGBColor(0x66, 0x70, 0x85)
    for n, (text, opts, ans, note) in enumerate(QUESTIONS, 1):
        lines = text.split('\n')
        p = doc.add_paragraph()
        p.add_run(f'Câu {n}: ').bold = True
        p.add_run(lines[0])
        for extra in lines[1:]: doc.add_paragraph(extra)
        if n == 2:      # ví dụ phương án ngắn viết trên một dòng, cách nhau bằng Tab
            doc.add_paragraph('\t'.join(f'{chr(65 + i)}. {o}' for i, o in enumerate(opts)))
        else:
            for i, o in enumerate(opts): doc.add_paragraph(f'{chr(65 + i)}. {o}')
        doc.add_paragraph().add_run(f'Đáp án: {ans}').bold = True
        if note: doc.add_paragraph(note).runs[0].italic = True
        doc.add_paragraph()
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
