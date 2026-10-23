"""Xuất bảng điểm ra file Excel."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import LOCAL_TZ, local

STATUS = {'submitted': 'Đã nộp', 'timeout': 'Hết giờ, tự nộp', 'doing': 'Đang làm'}
HEAD = ['STT', 'MSSV', 'Họ và tên', 'Lớp', 'Điểm (thang 10)', 'Số câu đúng', 'Tổng số câu', 'Bắt đầu', 'Nộp bài',
        'Thời gian làm (phút)', 'Trạng thái']
WIDTH = [6, 14, 28, 12, 15, 12, 12, 19, 19, 19, 16]


def _local(s):
    return local(s).replace(tzinfo=None) if s else None      # Excel không lưu múi giờ


def results_workbook(exam, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Bảng điểm'
    ws.append([f'BẢNG ĐIỂM: {exam["title"]}'])
    ws.append([f'Mã phòng thi: {exam["code"]}   ·   {exam["num_questions"]} câu   ·   {exam["duration_min"]} phút   ·   '
               f'xuất lúc {datetime.now(LOCAL_TZ).strftime("%H:%M %d/%m/%Y")}'])
    ws.append([])
    ws.append(HEAD)
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'].font = Font(italic=True, color='667085')
    thin = Side(style='thin', color='D0D5DD')
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=4, column=i)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='2558E8')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for n, r in enumerate(rows, 1):
        start, end = _local(r['started_at']), _local(r['submitted_at'])
        ws.append([n, r['student_code'], r['student_name'], r['class_name'] or '', r['score'], r['correct'], r['total'],
                   start, end, round((end - start).total_seconds() / 60, 1) if end else None, STATUS.get(r['status'], r['status'])])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if c.column in (8, 9): c.number_format = 'hh:mm:ss dd/mm/yyyy'
            if c.column == 5: c.number_format = '0.00'
    ws.freeze_panes = 'A5'
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
