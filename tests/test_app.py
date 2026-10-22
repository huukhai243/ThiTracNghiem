"""Chạy thử toàn bộ quy trình: giáo viên tạo đề từ file Word, sinh viên thi, chấm điểm, xuất Excel."""
import json
import os
import re
import shutil
import sqlite3
import tempfile
import unittest
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from quiz import create_app


def bank_docx(n=80):
    """Ngân hàng n câu; câu chia hết cho 5 có hai đáp án đúng (A, C), các câu khác đáp án B."""
    doc = Document()
    doc.add_paragraph('NGÂN HÀNG CÂU HỎI KIỂM THỬ')
    for i in range(1, n + 1):
        doc.add_paragraph(f'Câu {i}: Nội dung câu hỏi số {i}?')
        for k, letter in enumerate('ABCD'): doc.add_paragraph(f'{letter}. Phương án {letter} của câu {i}')
        doc.add_paragraph('Đáp án: A, C' if i % 5 == 0 else 'Đáp án: B')
    buf = BytesIO(); doc.save(buf); buf.seek(0)
    return buf


class Flow(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        os.environ['QUIZ_DATA_DIR'] = self.tmp
        self.app = create_app({'TESTING': True})
        self.teacher = self.client()
        self.cons = []

    def tearDown(self):
        for con in self.cons: con.close()
        os.environ.pop('QUIZ_DATA_DIR', None)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def client(self):
        c = self.app.test_client()
        with c.session_transaction() as s: s['csrf'] = 'tok'
        return c

    def db(self):
        con = sqlite3.connect(self.app.config['DATABASE']); con.row_factory = sqlite3.Row
        self.cons.append(con)
        return con

    def post(self, c, url, **data):
        data['csrf'] = 'tok'
        return c.post(url, data=data, content_type='multipart/form-data')

    def make_exam(self, **settings):
        self.assertEqual(self.teacher.get('/teacher/').status_code, 302)
        r = self.post(self.teacher, '/teacher/setup', username='gv01', password='matkhau1', password2='matkhau1')
        self.assertEqual(r.status_code, 302)
        with self.teacher.session_transaction() as s: s['csrf'] = 'tok'      # đăng nhập làm mới phiên
        form = dict(title='Kiểm tra giữa kỳ', num_questions='40', duration_min='30', shuffle_options='1', multi_mode='all',
                    start_at='', end_at='', file=(bank_docx(), 'bank.docx'))
        form.update(settings)
        r = self.post(self.teacher, '/teacher/exams/new', **form)
        self.assertEqual(r.status_code, 302, r.data.decode()[:500])
        exam = self.db().execute('SELECT * FROM exams').fetchone()
        self.assertEqual(self.post(self.teacher, f'/teacher/exams/{exam["id"]}/toggle').status_code, 302)
        return self.db().execute('SELECT * FROM exams').fetchone()

    def join(self, c, code, mssv, name='Nguyễn Văn A'):
        return self.post(c, '/join', code=code.lower(), name=name, student_code=mssv, class_name='K67')

    def correct_answers(self, attempt_id, right_count):
        """Đáp án cho form nộp bài: right_count câu đầu đúng, còn lại chọn sai."""
        con, form = self.db(), {}
        for it in con.execute('SELECT * FROM attempt_items WHERE attempt_id=? ORDER BY position', (attempt_id,)):
            opts = con.execute('SELECT id, is_correct FROM options WHERE question_id=?', (it['question_id'],)).fetchall()
            good = [o['id'] for o in opts if o['is_correct']]
            bad = [o['id'] for o in opts if not o['is_correct']][:1]
            form[f'q{it["position"]}'] = [str(x) for x in (good if it['position'] <= right_count else bad)]
        return form

    def test_full_flow(self):
        exam = self.make_exam()
        self.assertEqual((exam['num_questions'], exam['is_open']), (40, 1))
        self.assertEqual(self.db().execute('SELECT COUNT(*) FROM questions').fetchone()[0], 80)

        # sinh viên 1 vào thi: 40 câu, không lộ đáp án trong trang
        s1 = self.client()
        self.assertEqual(self.join(s1, exam['code'], 'sv001').status_code, 302)
        page = s1.get('/exam').data.decode()
        self.assertEqual(page.count('class="question"'), 40)
        self.assertNotIn('correct', page)
        self.assertEqual(page.count('type="checkbox"') > 0, True)            # câu nhiều đáp án dùng ô chọn nhiều
        a1 = self.db().execute("SELECT * FROM attempts WHERE student_code='SV001'").fetchone()
        self.assertEqual((a1['total'], a1['student_name'], a1['class_name']), (40, 'Nguyễn Văn A', 'K67'))

        # tự lưu một câu
        it = self.db().execute('SELECT * FROM attempt_items WHERE attempt_id=? AND position=1', (a1['id'],)).fetchone()
        oid = json.loads(it['option_ids'])[0]
        r = s1.post('/exam/save', json={'pos': 1, 'selected': [oid, 999999]}, headers={'X-CSRF-Token': 'tok'})
        self.assertTrue(r.get_json()['ok'])
        saved = json.loads(self.db().execute('SELECT selected FROM attempt_items WHERE attempt_id=? AND position=1',
                                             (a1['id'],)).fetchone()[0])
        self.assertEqual(saved, [oid])                                      # id lạ bị bỏ

        # sinh viên 2 nhận bộ câu hỏi khác
        s2 = self.client()
        self.join(s2, exam['code'], 'SV002', 'Trần Thị B')
        q1 = {r[0] for r in self.db().execute('SELECT question_id FROM attempt_items WHERE attempt_id=?', (a1['id'],))}
        a2 = self.db().execute("SELECT id FROM attempts WHERE student_code='SV002'").fetchone()[0]
        q2 = {r[0] for r in self.db().execute('SELECT question_id FROM attempt_items WHERE attempt_id=?', (a2,))}
        self.assertNotEqual(q1, q2)

        # cùng MSSV trên máy khác bị chặn
        other = self.client()
        r = self.join(other, exam['code'], 'SV001')
        self.assertEqual(r.status_code, 400)
        self.assertIn('máy khác', r.data.decode())

        # nộp bài: 30/40 câu đúng -> 7.5 điểm
        r = self.post(s1, '/exam/submit', auto='0', **self.correct_answers(a1['id'], 30))
        self.assertEqual(r.status_code, 302)
        a1 = self.db().execute('SELECT * FROM attempts WHERE id=?', (a1['id'],)).fetchone()
        self.assertEqual((a1['status'], a1['score'], a1['correct']), ('submitted', 7.5, 30))
        self.assertIn('Đã nộp bài', s1.get('/done').data.decode())
        self.assertEqual(s1.get('/exam').status_code, 302)                  # nộp rồi không vào lại được
        self.assertIn('đã nộp', self.join(self.client(), exam['code'], 'SV001').data.decode())

        # sinh viên 2 bỏ dở, quá giờ -> tự chấm khi giáo viên mở bảng điểm
        con = self.db()
        with con: con.execute("UPDATE attempts SET deadline='2000-01-01T00:00:00+00:00' WHERE id=?", (a2,))
        page = self.teacher.get(f'/teacher/exams/{exam["id"]}/results').data.decode()
        self.assertIn('7.50', page)
        self.assertIn('Hết giờ, tự nộp', page)
        a2row = self.db().execute('SELECT * FROM attempts WHERE id=?', (a2,)).fetchone()
        self.assertEqual((a2row['status'], a2row['score']), ('timeout', 0.0))

        # xem chi tiết bài làm và xuất Excel
        self.assertEqual(self.teacher.get(f'/teacher/exams/{exam["id"]}/results/{a1["id"]}').status_code, 200)
        r = self.teacher.get(f'/teacher/exams/{exam["id"]}/results.xlsx')
        ws = load_workbook(BytesIO(r.data)).active
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
        self.assertEqual([(x[1], x[2], x[4], x[5], x[10]) for x in rows],
                         [('SV001', 'Nguyễn Văn A', 7.5, 30, 'Đã nộp'), ('SV002', 'Trần Thị B', 0.0, 0, 'Hết giờ, tự nộp')])

        # đổi cách chấm câu nhiều đáp án -> chấm lại; không thay được ngân hàng khi đã có bài làm
        r = self.post(self.teacher, f'/teacher/exams/{exam["id"]}/settings', title='Kiểm tra giữa kỳ', num_questions='40',
                      duration_min='30', shuffle_options='1', multi_mode='partial', start_at='', end_at='')
        self.assertEqual(r.status_code, 302)
        self.post(self.teacher, f'/teacher/exams/{exam["id"]}/upload', file=(bank_docx(10), 'x.docx'))
        self.assertEqual(self.db().execute('SELECT COUNT(*) FROM questions').fetchone()[0], 80)

        # đóng phòng: không vào thi mới được
        self.post(self.teacher, f'/teacher/exams/{exam["id"]}/toggle')
        r = self.join(self.client(), exam['code'], 'SV003')
        self.assertIn('chưa mở', r.data.decode())

    def test_resume_on_another_device(self):
        exam = self.make_exam()
        a, b = self.client(), self.client()
        self.join(a, exam['code'], 'SV010')
        self.assertEqual(self.join(b, exam['code'], 'SV010').status_code, 400)
        aid = self.db().execute('SELECT id FROM attempts').fetchone()[0]
        self.post(self.teacher, f'/teacher/exams/{exam["id"]}/results/{aid}/resume')
        self.assertEqual(self.join(b, exam['code'], 'SV010').status_code, 302)
        self.assertEqual(b.get('/exam').status_code, 200)
        self.assertEqual(a.get('/exam').status_code, 302)                   # máy cũ mất quyền
        self.post(self.teacher, f'/teacher/exams/{exam["id"]}/results/{aid}/reset')
        self.assertEqual(self.db().execute('SELECT COUNT(*) FROM attempts').fetchone()[0], 0)

    def test_late_submit_only_keeps_saved_answers(self):
        exam = self.make_exam()
        s = self.client()
        self.join(s, exam['code'], 'SV020')
        aid = self.db().execute('SELECT id FROM attempts').fetchone()[0]
        con = self.db()
        with con: con.execute("UPDATE attempts SET deadline='2000-01-01T00:00:00+00:00' WHERE id=?", (aid,))
        self.post(s, '/exam/submit', auto='0', **self.correct_answers(aid, 40))
        row = self.db().execute('SELECT status, score FROM attempts').fetchone()
        self.assertEqual((row['status'], row['score']), ('timeout', 0.0))

    def test_csrf_and_login_required(self):
        c = self.app.test_client()
        self.assertEqual(c.post('/join', data={'code': 'X'}).status_code, 400)
        self.assertEqual(c.get('/teacher/exams/1').status_code, 302)
        r = c.get('/teacher/sample.docx')                                   # file mẫu tải được không cần đăng nhập
        self.assertEqual((r.status_code, r.data[:2]), (200, b'PK'))

    def test_bad_word_file(self):
        self.make_exam()
        r = self.post(self.teacher, '/teacher/exams/new', title='Lỗi', num_questions='40', duration_min='30', multi_mode='all',
                      file=(BytesIO(b'not a docx'), 'x.docx'))
        self.assertEqual(r.status_code, 400)
        self.assertIn('Không đọc được file', r.data.decode())
        empty = Document(); empty.add_paragraph('chỉ có chữ'); buf = BytesIO(); empty.save(buf); buf.seek(0)
        r = self.post(self.teacher, '/teacher/exams/new', title='Rỗng', num_questions='40', duration_min='30', multi_mode='all',
                      file=(buf, 'y.docx'))
        self.assertIn('Không tìm thấy câu hỏi nào', r.data.decode())


if __name__ == '__main__':
    unittest.main()
